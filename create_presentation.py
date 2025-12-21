#!/usr/bin/env python3
"""
Generate PowerPoint presentation for Date Fruit Harvesting System Improvements
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Title
    left = Inches(0.5)
    top = Inches(2.5)
    width = Inches(9)
    height = Inches(1.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    top = Inches(4)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(24)
    p.font.italic = True
    p.alignment = PP_ALIGN.CENTER

def add_content_slide(prs, title, content_lines):
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Title
    left = Inches(0.5)
    top = Inches(0.3)
    width = Inches(9)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True

    # Content
    top = Inches(1.3)
    height = Inches(5.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, line in enumerate(content_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.space_after = Pt(8)

def add_table_slide(prs, title, headers, rows):
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Title
    left = Inches(0.5)
    top = Inches(0.3)
    width = Inches(9)
    height = Inches(0.8)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True

    # Table
    num_rows = len(rows) + 1
    num_cols = len(headers)
    left = Inches(0.5)
    top = Inches(1.3)
    width = Inches(9)
    height = Inches(0.5 * num_rows)

    table = slide.shapes.add_table(num_rows, num_cols, left, top, width, height).table

    # Headers
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)

    # Data rows
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(value)
            cell.text_frame.paragraphs[0].font.size = Pt(12)

def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    add_title_slide(prs,
        "Autonomous Date Fruit Harvesting System",
        "Recent Improvements & Optimizations"
    )

    # Slide 2: Problems We Solved
    add_table_slide(prs, "Problems We Solved",
        ["Problem", "Impact"],
        [
            ["Orientation & direction disagreed", "Gripper approached from wrong angle"],
            ["Single-metric fruit selection", "Picked wrong fruit in cluttered scenes"],
            ["Excessive gripper rotation", "Slow, jerky movements"],
            ["Noisy depth readings", "Inaccurate grasp positions"],
            ["Unstable goal acceptance", "Pursued jittering/false detections"],
        ]
    )

    # Slide 3: Unified Orientation/Direction
    add_content_slide(prs, "Improvement 1: Unified Orientation/Direction", [
        "BEFORE:",
        "  Ellipse axis --> Orientation (shape-based)",
        "  Heatmap peak --> Direction (closest point)",
        "  Result: DISAGREED - Wrong approach angle",
        "",
        "AFTER:",
        "  Heatmap peak --> Direction --> BOTH orientation & approach",
        "  Result: ALWAYS AGREE - Correct gripper alignment",
        "",
        "File: date_v1.7.py (lines 910-951)"
    ])

    # Slide 4: Multi-Factor Scoring
    add_table_slide(prs, "Improvement 2: Multi-Factor Scoring System",
        ["Factor", "Weight", "What it measures"],
        [
            ["Distance", "30%", "Closer to gripper = better"],
            ["Visibility", "25%", "Unoccluded surface area"],
            ["Depth Quality", "20%", "Lower stereo noise"],
            ["Confidence", "15%", "YOLO detection score"],
            ["Ellipse", "10%", "Valid orientation fit"],
            ["Sticky Bonus", "+15%", "Same fruit (tracking persistence)"],
        ]
    )

    # Slide 5: SLERP Blending
    add_content_slide(prs, "Improvement 3: SLERP Quaternion Blending", [
        "PROBLEM: Gripper rotated unnecessarily between grasps",
        "",
        "SOLUTION:",
        "  1. Check if 180 deg flip is closer to current orientation",
        "  2. Pick closer option (gripper symmetry)",
        "  3. Blend: 70% current + 30% target (configurable)",
        "",
        "RESULT: Minimal rotation, smooth movements",
        "",
        "File: goals.py (lines 37-101)"
    ])

    # Slide 6: Direction-Biased Pre-Grasp
    add_content_slide(prs, "Improvement 4: Direction-Biased Pre-Grasp", [
        "BEFORE: Fixed approach offset",
        "",
        "AFTER: Smart blended direction",
        "",
        "Blended Direction =",
        "    60% x fruit_direction (from vision) +",
        "    40% x visibility_approach (camera-optimal)",
        "",
        "BENEFIT: Better camera view + respects fruit orientation",
        "",
        "File: goals.py (lines 386-401)"
    ])

    # Slide 7: Robust Depth Filtering
    add_content_slide(prs, "Improvement 5: Robust Depth Filtering", [
        "BEFORE: Raw depth values (noisy)",
        "",
        "AFTER:",
        "  - Median Z filtering (20-frame history)",
        "  - Outlier rejection (>8mm from median)",
        "  - Z tolerance check (>5cm = reject)",
        "",
        "RESULT: More reliable grasp positions",
        "",
        "File: goals.py (lines 125-171)"
    ])

    # Slide 8: Stable Goal Acceptance
    add_content_slide(prs, "Improvement 6: Stable Goal Acceptance", [
        "BEFORE: Accept any detection immediately",
        "",
        "AFTER:",
        "  - Require 2+ stable readings within 2cm",
        "  - Detect position jumps >10cm (new fruit)",
        "  - Use VOLATILE QoS (fresh data only)",
        "",
        "RESULT: No more chasing false detections",
        "",
        "File: goals.py (lines 218-280)"
    ])

    # Slide 9: System Architecture
    add_content_slide(prs, "System Architecture", [
        "PERCEPTION LAYER:",
        "  ZED Camera --> YOLO Detection --> Scoring & Selection",
        "                      |",
        "          +-----------+-----------+",
        "          v                       v",
        "  /external_goal_pose    /datefruit_direction",
        "",
        "PLANNING LAYER:",
        "  Goal Accept --> Direction Blend --> Orientation SLERP",
        "                      |",
        "                      v",
        "                cuRobo Motion Planning",
        "",
        "EXECUTION LAYER:",
        "  Approach --> Reacquire --> Final Grasp --> Dropoff"
    ])

    # Slide 10: Key Parameters
    add_table_slide(prs, "Key Parameters (Tunable)",
        ["Parameter", "Value", "Purpose"],
        [
            ["STICKY_BONUS", "0.15", "Fruit tracking persistence"],
            ["blend_weight", "0.0-1.0", "Orientation current vs target"],
            ["standoff", "8cm", "Pre-grasp offset distance"],
            ["z_tolerance", "5cm", "Reacquisition depth tolerance"],
            ["stable_count", "2", "Readings needed for acceptance"],
            ["publish_rate", "50Hz", "Goal update frequency"],
        ]
    )

    # Slide 11: Results
    add_table_slide(prs, "Results & Benefits",
        ["Metric", "Before", "After"],
        [
            ["Gripper alignment", "Often wrong", "Correct"],
            ["Fruit selection", "Unstable", "Consistent"],
            ["Rotation movement", "Excessive", "Minimal"],
            ["Depth accuracy", "Noisy", "Filtered"],
            ["False pursuits", "Frequent", "Rare"],
            ["Goal acceptance", "Immediate", "Stable"],
        ]
    )

    # Slide 12: Files Modified
    add_table_slide(prs, "Files Modified",
        ["File", "Key Changes"],
        [
            ["date_v1.7.py", "Unified orientation, scoring system, direction publishing"],
            ["goals.py", "SLERP blending, direction bias, depth filtering, reacquisition"],
            ["node.py", "Direction subscription, goal tracking, stability checks"],
        ]
    )

    # Slide 13: Future Improvements
    add_content_slide(prs, "Future Improvements", [
        "1. TRUE 3D DIRECTION",
        "   Use actual depth for approach vector (not just 2D + Z=0)",
        "",
        "2. SURFACE NORMAL ESTIMATION",
        "   Perpendicular approach to fruit surface",
        "",
        "3. DYNAMIC OBSTACLE AVOIDANCE",
        "   Real-time path replanning",
        "",
        "4. FORCE-BASED GRASP FEEDBACK",
        "   Adaptive grip strength based on sensor data"
    ])

    # Slide 14: Summary
    add_content_slide(prs, "Summary", [
        "KEY ACHIEVEMENTS:",
        "",
        "  - Unified orientation & direction from single source",
        "  - Robust multi-factor fruit scoring (5 weighted factors)",
        "  - Minimal-rotation gripper movements (SLERP blending)",
        "  - Noise-resistant depth processing (median filtering)",
        "  - Stable, reliable goal acceptance (stability checks)",
        "",
        "",
        "RESULT: More reliable autonomous date fruit harvesting"
    ])

    # Save
    output_path = "/home/datepalm2/Documents/v2/manipulatorsdatepalm/presentation_improvements.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    main()
