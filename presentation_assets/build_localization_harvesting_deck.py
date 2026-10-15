#!/usr/bin/env python3
"""Date Localization & Harvesting — 2-week technical update.

Two topics: (1) date localization — the direct TensorRT decoder that replaced
Ultralytics post-processing, the FP16 quantization/export pipeline, and a
Jetson comparison; (2) date harvesting outdoor deployment — field-model
replay results, direct-vs-Ultralytics agreement, the one/two-Jetson
comparison, and the four field-image results. Code and data are pulled
verbatim from the live repo (direct_trt_decoder.py, yolo_thread.py,
bin/quantinze.py, the engine's .build.json manifest); benchmark numbers not
independently re-derivable here (the SVO2 replay run, field-image detection
counts) are reported as provided. Local presentation file only -- not
published to GitLab.
"""
from pathlib import Path
import sys

sys.path.insert(0, "/tmp/datepalm_pydeps")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation_output" / "Date_Localization_Harvesting_Update_2026-09-03.pptx"
MEDIA = ROOT / "presentation_assets" / "farm_media" / "localization_evidence"

NAVY = RGBColor(14, 30, 48); BLUE = RGBColor(17, 96, 140); CYAN = RGBColor(20, 150, 160)
GREEN = RGBColor(45, 150, 95); ORANGE = RGBColor(214, 118, 40); RED = RGBColor(196, 66, 66)
WHITE = RGBColor(255, 255, 255); INK = RGBColor(28, 40, 52); MUTED = RGBColor(110, 126, 140)
PANEL = RGBColor(244, 247, 249); LINE = RGBColor(210, 219, 225)
CODE_BG = RGBColor(20, 26, 33); CODE_TXT = RGBColor(214, 224, 232)
CODE_OLD = RGBColor(230, 130, 120); CODE_NEW = RGBColor(120, 210, 160)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)


def set_bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def textbox(slide, text, x, y, w, h, size=13, color=INK, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos", italic=False):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    return b


def rrect(slide, x, y, w, h, fill, border=None, radius=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = border or fill; s.line.width = Pt(1.0)
    s.shadow.inherit = False
    return s


def title(slide, headline, kicker, dark=False):
    kc = CYAN if not dark else CYAN
    hc = WHITE if dark else NAVY
    textbox(slide, kicker.upper(), 0.55, 0.32, 12.2, 0.28, 10.5, kc, True)
    textbox(slide, headline, 0.52, 0.60, 12.3, 0.62, 23, hc, True)
    rrect(slide, 0.55, 1.24, 0.9, 0.045, CYAN, radius=False)


def footer(slide, n, total=7, note="Date Localization & Harvesting — 2-week update — 2026-09-03"):
    textbox(slide, note, 0.55, 7.16, 10.5, 0.2, 7.5, MUTED)
    textbox(slide, f"{n} / {total}", 12.15, 7.12, 0.7, 0.24, 9, CYAN, True, PP_ALIGN.RIGHT)


def slide(bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, bg); return s


def code_panel(slide, x, y, w, h, lines, label, label_color):
    rrect(slide, x, y, w, 0.34, label_color, radius=False)
    textbox(slide, label, x + 0.12, y, w - 0.24, 0.34, 11, WHITE, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    rrect(slide, x, y + 0.34, w, h - 0.34, CODE_BG, radius=False)
    b = slide.shapes.add_textbox(Inches(x + 0.12), Inches(y + 0.42), Inches(w - 0.24), Inches(h - 0.5))
    tf = b.text_frame; tf.clear(); tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.font.name = "Consolas"; p.space_after = Pt(1)
        r = p.add_run(); r.text = line
        r.font.size = Pt(9.5); r.font.name = "Consolas"; r.font.color.rgb = CODE_TXT


def metric(slide, value, label, x, y, w=2.2, accent=CYAN):
    rrect(slide, x, y, w, 1.02, PANEL, LINE)
    textbox(slide, value, x + .13, y + .10, w - .26, .46, 19, accent, True)
    textbox(slide, label, x + .13, y + .58, w - .26, .38, 9, MUTED, True)


def bullets(slide, items, x, y, w, h, size=12.5, spacing=7):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.clear(); tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.font.name = "Aptos"; p.space_after = Pt(spacing)
        head, _, rest = it.partition("|")
        r = p.add_run(); r.text = "▸  " + head.strip()
        r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = INK
        if rest:
            r2 = p.add_run(); r2.text = "  " + rest.strip()
            r2.font.size = Pt(size); r2.font.color.rgb = MUTED
    return b


def table(slide, x, y, w, headers, rows, col_w=None, row_h=0.36, header_h=0.34,
          size=9.5, header_size=9.5):
    n = len(headers)
    col_w = col_w or [w / n] * n
    cx = x
    rrect(slide, x, y, w, header_h, NAVY, radius=False)
    for i, htext in enumerate(headers):
        textbox(slide, htext, cx + 0.05, y, col_w[i] - 0.1, header_h, header_size, WHITE, True,
                PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
        cx += col_w[i]
    ry = y + header_h
    for ridx, row in enumerate(rows):
        rowbg = PANEL if ridx % 2 == 0 else WHITE
        rrect(slide, x, ry, w, row_h, rowbg, LINE, radius=False)
        cx = x
        for i, cell in enumerate(row):
            textbox(slide, str(cell), cx + 0.05, ry, col_w[i] - 0.1, row_h, size, INK, False,
                    PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
            cx += col_w[i]
        ry += row_h
    return ry


def img(slide, path, x, y, w, h=None):
    path = Path(path)
    if not path.exists():
        rrect(slide, x, y, w, h or w * .56, PANEL, LINE)
        textbox(slide, "(missing)", x, y + (h or w * .56) / 2 - .15, w, .3, 10, MUTED, True, PP_ALIGN.CENTER)
        return
    if h is None:
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


# ============================================================================
# SLIDE 1 -- Localization: Ultralytics post-processing replaced
# ============================================================================
s = slide()
title(s, "Date Localization: Replacing Ultralytics Post-Processing", "Part 1 · Detection pipeline")
textbox(s, "yolo_thread.py — the inference call, unchanged surrounding code, before vs. after", 0.55, 1.36, 12, 0.3, 11, MUTED)

code_panel(s, 0.55, 1.75, 6.05, 4.55, [
    "# BEFORE — Ultralytics YOLO.predict() orchestration",
    "det = self._model.predict(",
    "     img,",
    "     save=False,",
    "     retina_masks=False,",
    "     imgsz=self.img_size,",
    "     conf=self.conf_thres,",
    "     iou=0.3,",
    "     max_det=self.max_det,",
    "     device=device,",
    "     verbose=False,",
    "     classes=self._detect_class_ids or None,",
    " )[0]",
    "",
    "# TensorRT kernel: ~9-17ms",
    "# Total predict() wall time: ~80-150ms",
    "# (framework post-processing + CPU mask",
    "#  work dominates, not the TRT kernel)",
], "BEFORE — yolo_thread.py (Ultralytics path)", RED)

code_panel(s, 6.75, 1.75, 6.05, 4.55, [
    "# AFTER — DirectTensorRTSegmenter",
    "det = self._direct_model.predict(img)",
    "",
    "# direct_trt_decoder.py, predict():",
    "padded = self.letterbox(image=image)",
    "tensor = torch.from_numpy(...)",
    "raw = self.backend(tensor)  # AutoBackend",
    "",
    "rows = predictions[0]",
    "rows = rows[rows[:,4] > self.conf][:max_det]",
    "rows = rows[class filter]",
    "",
    "masks = ops.process_mask(",
    "    prototypes[0], rows[:,6:], rows[:,:4],",
    "    tensor.shape[2:], upsample=True)",
    "",
    "return Results(image, boxes=boxes, masks=masks)",
], "AFTER — direct_trt_decoder.py (82 lines)", GREEN)

textbox(s, "Same Ultralytics-compatible Results object returned either way — downstream depth/localization "
           "code is unchanged. Permanent fallback to the Ultralytics path if the direct decoder raises an exception "
           "(UR10E_DIRECT_TRT_DECODER=0 also restores it manually). Runtime selection: yolo_thread.py.",
        0.55, 6.42, 12.2, 0.7, 10.5, MUTED)
footer(s, 1)

# ============================================================================
# SLIDE 2 -- Localization: Quantization / engine export
# ============================================================================
s = slide()
title(s, "FP16 Quantization: Reproducible, Device-Locked Engine Export", "Part 1 · Engine build pipeline")

bullets(s, [
    "Fixed shape, no NMS in-graph | imgsz 1024 (or 1248), batch 1, dynamic=False, simplify=True, nms=False",
    "FP16 by default | half=True; INT8 only when explicitly requested with a representative calibration dataset",
    "Built on the exact target Jetson | script refuses to run unless /etc/nv_tegra_release exists on the host",
    "MAXN + jetson_clocks | doesn't change FP16 weights numerically — stabilizes TensorRT's tactic "
    "benchmarking during the build so the selected kernels aren't chosen under throttled clocks",
    "Refuses to silently overwrite | an existing .engine at the destination name",
    "Embedded verification | size, batch, precision, dynamic-shape and NMS settings checked against the export request",
], 0.55, 1.75, 6.5, 4.3, size=12, spacing=10)

textbox(s, ".build.json manifest (real, this engine)", 7.35, 1.75, 5.3, 0.3, 12, NAVY, True)
code_panel(s, 7.35, 2.08, 5.35, 3.55, [
    "\"device\": \"Orin\",",
    "\"compute_capability\": [8, 7],",
    "\"cuda\": \"12.2\",  \"tensorrt\": \"8.6.2\",",
    "\"ultralytics\": \"8.4.6\",",
    "\"nvpmodel\": \"NV Power Mode: MAXN\",",
    "\"export_settings\": {",
    "  \"imgsz\": 1024, \"batch\": 1,",
    "  \"dynamic\": false, \"simplify\": true,",
    "  \"nms\": false, \"half\": true, \"int8\": false",
    "},",
    "\"source_sha256\": \"6cd374121aa...\",",
    "\"engine_sha256\": \"3ce6602b8d8...\",",
], "30_Aug_26_kaust_farm_1024_fp16.build.json", BLUE)

metric(s, "1024²", "fixed input shape", 0.55, 6.15, 2.05, CYAN)
metric(s, "FP16", "default precision", 2.72, 6.15, 2.05, GREEN)
metric(s, "sha256", "source + engine hashed", 4.89, 6.15, 2.35, ORANGE)
textbox(s, "How this improved results: a fixed shape and no in-graph NMS let TensorRT pick simpler, faster kernels "
           "and removed dynamic-shape overhead; per-Jetson builds under MAXN mean the deployed kernels were "
           "actually benchmarked at full clock, not chosen under an arbitrary throttled state.",
        7.35, 5.75, 5.35, 1.1, 10, MUTED)
footer(s, 2)

# ============================================================================
# SLIDE 3 -- Localization: One vs Two Jetsons
# ============================================================================
s = slide()
title(s, "One Jetson vs. Two: Localization Reliability", "Part 1 · Deployment topology")

hdrs = ["Configuration", "Observed behavior", "Advantage", "Main limitation"]
rows = [
    ["One Jetson, raw YOLO only", "~21–30 FPS (earlier raw tests)",
     "No ROS/network transport; lowest latency", "No depth, tracking, viz, or cuRobo"],
    ["One Jetson, full harvesting", "~8–12 YOLO FPS; camera loop ~27–30 FPS after decoupling",
     "Shared-memory image flow; simplest, most reliable", "YOLO/cuRobo contend for GPU/CPU/memory"],
    ["Two Jetsons, tested split", "~8 YOLO FPS, 10–12 loop FPS (2 dates observed)",
     "Can isolate cuRobo from inference, in principle", "DDS/QoS/MTU setup, 30W/8-core vision Jetson erased the gain"],
]
col_w = [2.55, 3.7, 3.35, 3.0]
table(s, 0.55, 1.75, 12.6, hdrs, rows, col_w=col_w, row_h=0.62, size=9.5)

textbox(s, "Localization angle: the second Jetson didn't change what the model detects — it changed how "
           "reliably and how often a fresh detection reaches the planner. A slower link/QoS mismatch means "
           "localization estimates the arm acts on can be stale, which matters more for grasp accuracy than raw FPS alone.",
        0.55, 4.05, 12.2, 0.7, 11, MUTED)

metric(s, "Single box", "current recommendation", 0.55, 5.0, 3.0, GREEN)
metric(s, "1024 · FP16", "Jetson-built + direct decoder", 3.75, 5.0, 3.35, CYAN)
metric(s, "Explicit handoff", "YOLO ↔ cuRobo GPU sharing", 7.3, 5.0, 3.0, ORANGE)
textbox(s, "Keep the two-Jetson design experimental until it transmits compact detections/3D targets "
           "(not full-rate raw/display/depth images), uses matched QoS and fixed wired addressing, and is "
           "benchmarked with both Jetsons verified in MAXN clocks.",
        0.55, 6.05, 12.2, 0.7, 10.5, MUTED, italic=True)
footer(s, 3)

# ============================================================================
# SLIDE 4 -- Harvesting: Field-model replay
# ============================================================================
s = slide()
title(s, "Field-Model Replay: SN57931814, 404 Frames", "Part 2 · Outdoor deployment — 2026-09-03")

textbox(s, "HD1080_SN57931814_03-01-29.svo2 replayed through 30_Aug_26_kaust_farm_1024_fp16.engine · "
           "conf=0.10 · max_det=10/frame. Full annotated video (mask + box + confidence + live count overlay), not selected stills.",
        0.55, 1.36, 12.2, 0.4, 10.5, MUTED)

hdrs = ["Recording", "Detected", "Mean/frame", "Range", "Transitions", "Mean conf.", "Direct TRT FPS*"]
rows = [["SN57931814", "404/404 (100%)", "9.82", "1–10", "10/403", "0.909", "36.9"]]
table(s, 0.55, 1.9, 12.6, hdrs, rows, col_w=[2.1, 1.9, 1.7, 1.4, 1.6, 1.7, 2.2], row_h=0.38)

textbox(s, "*Isolated replay throughput on this machine — excludes live ZED depth, tracking, ROS viz, "
           "cuRobo and network transport; not expected full-harvesting FPS. TensorRT flagged that this engine "
           "ran on a different device model than it was built on — deployment performance must be re-measured on the target Jetson.",
        0.55, 2.66, 12.2, 0.5, 9, MUTED, italic=True)

textbox(s, "Direct decoder vs. Ultralytics — every 5th frame, both paths", 0.55, 3.28, 8, 0.28, 12.5, NAVY, True)
hdrs2 = ["Recording", "Samples", "Exact count agreement", "Count MAE", "Direct mean", "Ultralytics mean", "Speed-up"]
rows2 = [["SN57931814", "81", "81/81", "0.00", "27.7 ms", "27.7 ms", "1.00×"]]
table(s, 0.55, 3.6, 12.6, hdrs2, rows2, col_w=[1.9, 1.4, 2.5, 1.5, 1.8, 1.9, 1.6], row_h=0.38)

bullets(s, [
    "Output-preserving | detection count matches exactly on sampled frames (81/81); speed benefit depends on scene/output load",
    "Mask reconstruction dominates | on this dense sequence, removing most of the wrapper-level speed gain — "
    "still need mask-IoU / box-coordinate comparison before claiming full output equivalence",
    "max_det=10 ceiling hit | in 392/404 frames — the model is consistently finding dates, but this replay "
    "can't measure total recall in dense bunches; a future run needs a higher cap + ground truth",
    "Field-domain risks in the scenes | direct-sun saturation, backlighting, stem overlap, dense occlusion, "
    "ripeness/color variation, red gripper fingertips in-frame",
], 0.55, 4.45, 12.2, 2.3, size=10.5, spacing=6)
footer(s, 4)

# ============================================================================
# SLIDE 5 -- Harvesting: One vs Two Jetsons (operational)
# ============================================================================
s = slide()
title(s, "One Jetson vs. Two: Harvesting Throughput", "Part 2 · Outdoor deployment — operational")

hdrs = ["Configuration", "Observed behavior", "Advantage", "Main limitation"]
rows = [
    ["One Jetson, raw YOLO only", "~21–30 FPS (earlier raw tests)",
     "No ROS/network transport; lowest latency", "No depth, tracking, viz, or cuRobo"],
    ["One Jetson, full harvesting", "~8–12 YOLO FPS; camera loop ~27–30 FPS after decoupling",
     "Shared-memory image flow; simplest, most reliable", "YOLO/cuRobo contend for GPU/CPU/memory"],
    ["Two Jetsons, tested split", "~8 YOLO FPS, 10–12 loop FPS (2 dates observed)",
     "Can isolate cuRobo from inference, in principle",
     "ROS2/DDS serialization, QoS mismatch, Ethernet MTU, large image/depth topic copies, "
     "30W/8-core vision Jetson (GPU capped 612MHz) erased the compute benefit"],
]
col_w = [2.4, 3.5, 3.0, 3.7]
table(s, 0.55, 1.75, 12.6, hdrs, rows, col_w=col_w, row_h=0.68, size=9.5)

textbox(s, "The second Jetson is not automatically faster. Compute was separated, but a new transport "
           "pipeline was introduced: acquisition/inference results cross ROS2/DDS, high-bandwidth image "
           "topics are serialized/copied, QoS must match, and the tested vision Jetson initially ran "
           "MODE_30W with only CPUs 0–7 online and GPU capped at 612MHz. The remote /vision/display stream "
           "also stopped or became incompatible when link/QoS settings changed.",
        0.55, 4.15, 12.2, 1.0, 11, INK)

rrect(s, 0.55, 5.55, 12.2, 1.45, PANEL, LINE)
textbox(s, "CURRENT RECOMMENDATION", 0.75, 5.67, 5, 0.26, 10.5, CYAN, True)
textbox(s, "Use the stronger Jetson as a single-box baseline: Jetson-built 1024 FP16 engine, direct decoder, "
           "latest-frame/drop-old queue, explicit YOLO↔cuRobo GPU handoff. Keep the two-Jetson design "
           "experimental until it transmits compact detections/3D targets (not full-rate raw/display/depth "
           "images), matched QoS, fixed wired addressing, and is benchmarked with both Jetsons in verified MAXN clocks.",
        0.75, 5.95, 11.8, 1.0, 10.5, INK)
footer(s, 5)

# ============================================================================
# SLIDE 6 -- Harvesting: Field image results
# ============================================================================
s = slide()
title(s, "Exact Uploaded Field-Image Results", "Part 2 · 30_Aug_26_kaust_farm_1024_fp16.engine · conf=0.10, max_det=30")

imgs = [
    ("camera_raw_20260902_165800.png", "19 detections · conf 0.127–0.916",
     "Dense/occluded — broad score spread"),
    ("camera_raw_20260831_172917.png", "13 detections · conf 0.206–0.960",
     "First of near-identical temporal pair"),
    ("camera_raw_20260831_172918.png", "13 detections · conf 0.179–0.954",
     "Count stable; weakest scores vary"),
    ("camera_raw_19700101_033237.png", "30 detections · conf 0.875–0.952",
     "Hit max_det=30 cap — lower bound only"),
]
positions = [(0.55, 1.75), (3.55, 1.75), (6.55, 1.75), (9.55, 1.75)]
for (fname, stat, note), (x, y) in zip(imgs, positions):
    img(s, MEDIA / fname, x, y, 2.85, 1.9)
    textbox(s, stat, x, y + 1.94, 2.85, 0.3, 10, NAVY, True)
    textbox(s, note, x, y + 2.24, 2.85, 0.5, 9, MUTED)

textbox(s, "The two consecutive Aug-31 frames show a stable count with strong high-confidence detections "
           "despite minor frame-to-frame score movement. The Sep-2 scene includes many partially occluded, "
           "lower-confidence candidates. The dense 1970-timestamp image saturates the evaluation cap — recall "
           "must be re-measured with a higher cap and ground-truth labels.",
        0.55, 4.6, 12.6, 0.85, 11, INK)
textbox(s, "Timing for the first still includes TensorRT warm-up; the other three direct-decoder calls were "
           "~24–31ms and are not full-pipeline latency.",
        0.55, 5.5, 12.6, 0.35, 9.5, MUTED, italic=True)
footer(s, 6)

# ============================================================================
# SLIDE 7 -- Remaining validation / next steps (shared)
# ============================================================================
s = slide()
title(s, "Remaining Validation", "Shared next steps — both parts")
bullets(s, [
    "Annotate the retained recording | report precision, recall, mAP, mask IoU, false-positive type, missed-date type",
    "Repeat without max_det=10 | measure dense-bunch recall without the evaluation ceiling",
    "Compare direct/Ultralytics boxes and masks | not only detection counts — needed before claiming full output equivalence",
    "Synchronized one- vs two-Jetson trials | identical SVO frames, fixed power mode/clocks/model hash/confidence/max_det",
    "End-to-end latency | camera timestamp → accepted robot goal, not just model throughput",
], 0.7, 1.9, 11.9, 4.2, size=15, spacing=16)
footer(s, 7)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"saved: {OUT}  ({len(prs.slides._sldIdLst)} slides)")
